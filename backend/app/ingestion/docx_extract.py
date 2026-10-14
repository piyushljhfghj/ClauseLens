"""DOCX text extraction.

The body is walked in document order (paragraphs, tables, content controls). Paragraph text is
read from the runs ourselves rather than via python-docx's `Paragraph.text`, so that:
- tracked insertions are included and tracked deletions excluded (the document as it would read
  with all changes accepted),
- tab-stop definitions in paragraph properties are not mistaken for text,
- Word auto-numbering is resolved and prefixed ("8.2 The Company may...").

Word documents have no fixed pages, so blocks carry no page numbers. Known gaps: text boxes,
footnotes and headers/footers are not extracted.
"""

import io
import re
from collections.abc import Callable

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.ingestion.docx_numbering import NumberingResolver
from app.ingestion.model import (
    CanonicalTextBuilder,
    ExtractedDocument,
    SourceBlock,
    SourceLine,
    SourceWord,
)

_EXCLUDED_ANCESTORS = {qn("w:del"), qn("w:moveFrom"), qn("w:txbxContent")}
_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.I)
_WHITESPACE = re.compile(r"\s+")

DocxProgress = Callable[[float], None]


def _is_excluded(element, stop) -> bool:  # type: ignore[no-untyped-def]
    parent = element.getparent()
    while parent is not None and parent is not stop:
        if parent.tag in _EXCLUDED_ANCESTORS:
            return True
        parent = parent.getparent()
    return False


def paragraph_text(p_element) -> str:  # type: ignore[no-untyped-def]
    parts: list[str] = []
    for run in p_element.iter(qn("w:r")):
        if _is_excluded(run, p_element):
            continue
        for child in run:
            tag = child.tag
            if tag == qn("w:t"):
                parts.append(child.text or "")
            elif tag == qn("w:tab"):
                parts.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
            elif tag == qn("w:noBreakHyphen"):
                parts.append("-")
    return "".join(parts)


def _lines(text: str) -> list[SourceLine]:
    lines = []
    for raw in text.split("\n"):
        words = [SourceWord(w, 0, 0, 0, 0) for w in _WHITESPACE.split(raw.strip()) if w]
        if words:
            lines.append(SourceLine(words=words))
    return lines


def _style_bold(style) -> bool:  # type: ignore[no-untyped-def]
    while style is not None:
        if style.font.bold is not None:
            return bool(style.font.bold)
        style = style.base_style
    return False


def is_all_bold(paragraph: Paragraph) -> bool:
    """True if every run with visible text is bold, directly or via its character/paragraph
    style. Real-world headings are often just bold "Normal" paragraphs, not Heading styles."""
    paragraph_bold = _style_bold(paragraph.style)
    visible = [run for run in paragraph.runs if run.text.strip()]
    if not visible:
        return False
    for run in visible:
        bold = run.bold
        if bold is None and run.style is not None:
            bold = run.style.font.bold
        if bold is None:
            bold = paragraph_bold
        if not bold:
            return False
    return True


def _heading_level(paragraph: Paragraph) -> int | None:
    outline = paragraph._p.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
    if outline is not None and outline.get(qn("w:val"), "").isdigit():
        level = int(outline.get(qn("w:val"))) + 1
        return level if level <= 9 else None  # 9 = "body text"
    style = paragraph.style
    while style is not None:
        name = style.name or ""
        if name.lower() == "title":
            return 1
        if match := _HEADING_STYLE.match(name):
            return int(match.group(1))
        style = style.base_style
    return None


def extract_docx(data: bytes, *, on_progress: DocxProgress | None = None) -> ExtractedDocument:
    document = docx.Document(io.BytesIO(data))
    numbering = NumberingResolver(document)
    builder = CanonicalTextBuilder()
    body = document.element.body

    def add_paragraph(p_element) -> None:  # type: ignore[no-untyped-def]
        paragraph = Paragraph(p_element, document)
        label, _ = numbering.label_for(p_element, paragraph.style)
        text = paragraph_text(p_element)
        if not text.strip():
            return
        if label:
            text = f"{label} {text}"
        lines = _lines(text)
        bold = is_all_bold(paragraph)
        for line in lines:
            line.bold = bold
        builder.add_block(
            SourceBlock(
                lines=lines,
                heading_level=_heading_level(paragraph),
                numbering_label=label,
            )
        )

    def add_table(tbl_element) -> None:  # type: ignore[no-untyped-def]
        for row in tbl_element.iter(qn("w:tr")):
            if _nearest_table(row) is not tbl_element:
                continue  # nested-table rows: their text is already read via the outer cell
            lines: list[SourceLine] = []
            for cell in row.findall(qn("w:tc")):
                for p in cell.iter(qn("w:p")):
                    lines.extend(_lines(paragraph_text(p)))
            if lines:
                builder.add_block(SourceBlock(lines=lines, kind="table_row"))

    def walk(container) -> None:  # type: ignore[no-untyped-def]
        children = list(container)
        for i, child in enumerate(children):
            if child.tag == qn("w:p"):
                add_paragraph(child)
            elif child.tag == qn("w:tbl"):
                add_table(child)
            elif child.tag == qn("w:sdt"):  # content control: walk its content
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    walk(content)
            if on_progress and container is body and i % 50 == 0:
                on_progress(i / max(len(children), 1))

    walk(body)
    if on_progress:
        on_progress(1.0)
    return ExtractedDocument(canonical_text=builder.build(), blocks=builder.blocks, pages=[])


def _nearest_table(element):  # type: ignore[no-untyped-def]
    parent = element.getparent()
    while parent is not None and parent.tag != qn("w:tbl"):
        parent = parent.getparent()
    return parent
