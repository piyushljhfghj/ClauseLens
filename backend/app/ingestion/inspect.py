"""Document inspection: open the file, count pages, and find out whether it has usable text.

Runs in a worker thread (PyMuPDF / python-docx are blocking). Problems with the *file itself*
raise `ProcessingFailure`, which is permanent: retrying will not fix a password-protected or
scanned PDF, so the document is marked failed with a specific, user-facing reason.
"""

import io
import logging
import zipfile
from collections.abc import Callable
from dataclasses import dataclass

import docx
import pymupdf
from docx.opc.exceptions import PackageNotFoundError

logger = logging.getLogger(__name__)

# A page with fewer extractable characters than this is treated as having no text layer.
MIN_TEXT_CHARS_PER_PAGE = 20

ProgressCallback = Callable[[float], None]  # fraction 0.0-1.0 of this step


class ProcessingFailure(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class InspectionResult:
    page_count: int | None  # None for DOCX: Word documents have no fixed pages
    pages_without_text: int | None


def inspect_pdf(
    data: bytes, *, max_pages: int, on_progress: ProgressCallback | None = None
) -> InspectionResult:
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ProcessingFailure(
            "CORRUPT_FILE", "This PDF could not be opened. The file may be damaged."
        ) from exc

    with document:
        if document.needs_pass:
            raise ProcessingFailure(
                "ENCRYPTED_PDF",
                "This PDF is password-protected. Remove the password and upload it again.",
            )
        if document.is_repaired:
            # MuPDF rebuilt a damaged file; unrecoverable pages surface as pages without text.
            logger.warning("damaged PDF was repaired on open")
        page_count = document.page_count
        if page_count == 0:
            raise ProcessingFailure("EMPTY_DOCUMENT", "This PDF has no pages.")
        if page_count > max_pages:
            raise ProcessingFailure(
                "TOO_MANY_PAGES",
                f"This PDF has {page_count} pages; the maximum supported is {max_pages}.",
            )

        pages_without_text = 0
        try:
            for index, page in enumerate(document):
                if len(page.get_text("text").strip()) < MIN_TEXT_CHARS_PER_PAGE:
                    pages_without_text += 1
                if on_progress:
                    on_progress((index + 1) / page_count)
        except Exception as exc:
            raise ProcessingFailure(
                "CORRUPT_FILE", "Some pages of this PDF could not be read. The file may be damaged."
            ) from exc

    if pages_without_text == page_count:
        raise ProcessingFailure(
            "SCANNED_NO_TEXT",
            "This PDF has no selectable text. It appears to be a scanned image, so its "
            "contents cannot be read. Upload a text-based PDF, or run it through OCR first.",
        )
    return InspectionResult(page_count=page_count, pages_without_text=pages_without_text)


def inspect_docx(data: bytes, *, on_progress: ProgressCallback | None = None) -> InspectionResult:
    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ProcessingFailure(
            "CORRUPT_FILE", "This Word document could not be opened. The file may be damaged."
        ) from exc
    except Exception as exc:  # malformed XML inside the package, etc.
        raise ProcessingFailure(
            "CORRUPT_FILE", "This Word document could not be read. The file may be damaged."
        ) from exc

    has_text = any(p.text.strip() for p in document.paragraphs) or any(
        cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells
    )
    if on_progress:
        on_progress(1.0)
    if not has_text:
        raise ProcessingFailure("EMPTY_DOCUMENT", "This Word document contains no text.")
    return InspectionResult(page_count=None, pages_without_text=None)


def inspect(
    file_type: str, data: bytes, *, max_pages: int, on_progress: ProgressCallback | None = None
) -> InspectionResult:
    if file_type == "pdf":
        return inspect_pdf(data, max_pages=max_pages, on_progress=on_progress)
    if file_type == "docx":
        return inspect_docx(data, on_progress=on_progress)
    raise ValueError(f"unknown file type: {file_type}")
