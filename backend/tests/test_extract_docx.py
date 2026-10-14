"""DOCX extraction: auto-numbering, tracked changes, tables, headings."""

import io
import re

import docx
import pytest
from docx.oxml import parse_xml

from app.ingestion.docx_extract import extract_docx, paragraph_text
from app.ingestion.docx_numbering import format_number
from tests.fixtures import factory


def collapse(text: str) -> str:
    return re.sub(r"\s+", " ", text)


@pytest.fixture(scope="module")
def numbered():  # type: ignore[no-untyped-def]
    return extract_docx(factory.numbered_docx())


def test_auto_numbering_is_resolved_into_text(numbered) -> None:  # type: ignore[no-untyped-def]
    text = collapse(numbered.canonical_text)
    for expected in [
        "1. DEFINITIONS",
        "1.1 “Services” means",
        "2. TERMINATION",
        "2.1 Either party may terminate",
        "2.2 The Company may terminate immediately",
        "(a) breaches any material term",
        "(b) becomes insolvent",
        "3. LIABILITY",
        "3.1 The total liability",
    ]:
        assert expected in text, expected


def test_tracked_changes_read_as_accepted(numbered) -> None:  # type: ignore[no-untyped-def]
    text = collapse(numbered.canonical_text)
    assert "Payment is due within forty-five days." in text
    assert "within thirty days" not in text and "forty-fivethirty" not in text


def test_tables_become_row_blocks(numbered) -> None:  # type: ignore[no-untyped-def]
    rows = [
        numbered.canonical_text[b.char_start : b.char_end]
        for b in numbered.blocks
        if b.kind == "table_row"
    ]
    assert [collapse(r) for r in rows] == ["Fee AED 100,000", "Term 24 months"]


def test_heading_style_is_captured(numbered) -> None:  # type: ignore[no-untyped-def]
    first = numbered.blocks[0]
    assert numbered.canonical_text[first.char_start : first.char_end] == "MASTER SERVICES AGREEMENT"
    assert first.heading_level == 1


def test_docx_has_no_pages_and_blocks_cover_text(numbered) -> None:  # type: ignore[no-untyped-def]
    assert numbered.pages == []
    assert all(b.page_start is None for b in numbered.blocks)
    for block in numbered.blocks:
        assert numbered.canonical_text[block.char_start : block.char_end].strip()


def test_tab_stop_definitions_are_not_text() -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Clause text")
    paragraph._p.get_or_add_pPr().append(
        parse_xml(
            '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:tab w:val="left" w:pos="720"/></w:tabs>'
        )
    )
    assert paragraph_text(paragraph._p) == "Clause text"


def test_plain_document_without_numbering() -> None:
    doc = extract_docx(factory.docx_file(["First paragraph.", "Second paragraph."]))
    assert doc.canonical_text == "First paragraph.\n\nSecond paragraph."


def test_numbering_continues_and_restarts_like_word() -> None:
    document = docx.Document(io.BytesIO(factory.numbered_docx()))
    text = collapse(extract_docx(factory.numbered_docx()).canonical_text)
    # Level-2 numbering restarts under each new top-level clause: 2.1 after 1.1, not 1.2.
    assert text.index("1.1 ") < text.index("2.1 ") < text.index("3.1 ")
    assert "1.2 " not in text
    assert document is not None


@pytest.mark.parametrize(
    ("n", "fmt", "expected"),
    [
        (3, "decimal", "3"),
        (3, "lowerLetter", "c"),
        (28, "lowerLetter", "bb"),
        (4, "upperRoman", "IV"),
        (9, "lowerRoman", "ix"),
        (5, "decimalZero", "05"),
        (2, "unknownFormat", "2"),
    ],
)
def test_format_number(n: int, fmt: str, expected: str) -> None:
    assert format_number(n, fmt) == expected
