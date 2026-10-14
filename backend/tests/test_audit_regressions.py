"""Regression tests for defects confirmed in the Phase 3 correctness audit.

Each test names the defect it guards against. See the audit report for the original evidence.
"""

import re

import docx
import pymupdf
import pytest

from app.ingestion import pdf_extract
from app.ingestion.docx_extract import extract_docx, is_all_bold
from app.ingestion.model import CanonicalTextBuilder, SourceBlock, SourceLine, SourceWord
from app.ingestion.ocr import tesseract_version
from app.ingestion.pdf_extract import OcrSettings, extract_pdf
from app.ingestion.structure import detect_sections, match_numbered_heading, section_preview
from tests.fixtures import factory

OCR_ON = OcrSettings(enabled=True, language="eng", dpi=300)
OCR_OFF = OcrSettings(enabled=False, language="eng", dpi=300)
EN_DASH, ELLIPSIS = chr(0x2013), chr(0x2026)
needs_tesseract = pytest.mark.skipif(tesseract_version() is None, reason="Tesseract not installed")


def collapse(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def _blocks(*items: tuple[str, bool], ocr: bool = False):  # type: ignore[no-untyped-def]
    builder = CanonicalTextBuilder()
    for text, bold in items:
        line = SourceLine(
            [SourceWord(w, 0, 0, 1, 1) for w in text.split()], font_size=10, bold=bold
        )
        block = builder.add_block(SourceBlock(lines=[line]))
        assert block is not None
        block.ocr = ocr
    return builder.build(), builder.blocks


# --- D1: numbered questions shown as bare numbers ------------------------------------------------


@pytest.mark.parametrize(
    ("text", "title"),
    [
        (
            "4. How did this course help you grow? Explain in the context of your strengths.",
            "How did this course help you grow?",
        ),
        (
            "1. What dimensions of your personality need to be developed to make you a "
            "successful professional?",
            "What dimensions of your personality need to be developed to make you a "
            "successful professional?",
        ),
        # Unchanged behaviour: a clause sentence is not a title; a short lead phrase is.
        ("8.2 The Company may terminate this Agreement upon notice.", None),
        (
            "8.2 Termination for Convenience. The Company may terminate.",
            "Termination for Convenience",
        ),
    ],
)
def test_d1_titles_of_numbered_questions(text: str, title: str | None) -> None:
    match = match_numbered_heading(text)
    assert match is not None and match.title == title


@pytest.mark.parametrize(
    ("text", "has_number", "preview"),
    [
        (
            "5. Enlist your semester achievements.\n\nAchievements: ...",
            True,
            "Enlist your semester achievements.",
        ),
        (f"Article IV {EN_DASH} Payment of the fees", True, "Payment of the fees"),
        ("Section 12.3: Governing Law", True, "Governing Law"),
        ("MASTER SERVICES AGREEMENT", False, "MASTER SERVICES AGREEMENT"),
    ],
)
def test_d1_section_preview_for_untitled_sections(
    text: str, has_number: bool, preview: str
) -> None:
    assert section_preview(text, has_number=has_number) == preview


def test_d1_long_preview_is_truncated_at_a_word_boundary() -> None:
    preview = section_preview(
        "8.2 " + "The Company may terminate this Agreement " * 5, has_number=True
    )
    assert preview.endswith(ELLIPSIS) and len(preview) <= 81
    assert not preview[:-1].endswith(" ")


# --- D2: DOCX bold headings ignored --------------------------------------------------------------


def test_d2_bold_normal_paragraphs_are_detected_as_headings() -> None:
    doc = extract_docx(factory.bold_headings_docx())
    sections = detect_sections(doc.blocks, doc.canonical_text)
    assert [(s.number, s.title) for s in sections] == [
        (None, "Personal Information"),
        (None, "About Yourself"),
        ("1", "What are your strengths and weaknesses?"),
        (None, "My Learning Experiences"),
        ("1", "How did this course help you grow?"),
    ]
    # The second question list starts its own section instead of being merged into the first.
    second = sections[-1]
    assert doc.canonical_text[second.char_start :].startswith("1. How did this course")


def test_d2_is_all_bold() -> None:
    document = docx.Document()
    full = document.add_paragraph()
    full.add_run("Personal Information").bold = True
    partial = document.add_paragraph()
    partial.add_run("Name: ").bold = True
    partial.add_run("A. Student")
    styled = document.add_paragraph("Styled", style="Heading 1")  # bold via paragraph style
    plain = document.add_paragraph("Plain text")
    assert [is_all_bold(p) for p in (full, partial, styled, plain)] == [True, False, True, False]


def test_d2_title_block_of_emphasised_lines_is_not_a_series_of_sections() -> None:
    text, blocks = _blocks(
        ("PERSONALITY AND HUMAN VALUES", True),
        ("JOURNAL FOR SUCCESS", True),
        ("Amity University", True),
        ("Personal Information", True),
        ("Name of the student: A. Student", False),
        ("About Yourself", True),
        ("1. What are your strengths?", False),
        ("I am curious.", False),
    )
    sections = detect_sections(blocks, text)
    assert [(s.number, s.title) for s in sections] == [
        (None, "Personal Information"),
        (None, "About Yourself"),
        ("1", "What are your strengths?"),
    ]


def test_d2_numbered_headings_are_never_demoted_by_the_title_block_rule() -> None:
    text, blocks = _blocks(
        ("ARTICLE 1 DEFINITIONS", True), ("1.1 Scope", True), ("Body text.", False)
    )
    assert [s.number for s in detect_sections(blocks, text)] == ["Article 1", "1.1"]


# --- D3: OCR noise promoted to headings -----------------------------------------------------------


def test_d3_ocr_text_does_not_produce_emphasis_headings() -> None:
    text, blocks = _blocks(
        ("AAA CY", False),
        ("VFOAAAK", False),
        ("8. TERMINATION", False),
        ("The parties agree.", False),
        ocr=True,
    )
    sections = detect_sections(blocks, text)
    assert [(s.number, s.title) for s in sections] == [("8", "TERMINATION")]  # numbering trusted


# --- D4: stamped scan silently not OCR'd ---------------------------------------------------------


@needs_tesseract
def test_d4_stamped_scan_is_ocrd() -> None:
    doc = extract_pdf(factory.stamped_scan_pdf(), ocr=OCR_ON)
    assert doc.pages[0].extraction == "ocr"
    assert "terminate this Agreement" in collapse(doc.canonical_text)
    for start, end, *_ in doc.pages[0].words:
        assert doc.canonical_text[int(start) : int(end)].strip()


def test_d4_stamped_scan_without_ocr_is_reported_unreadable() -> None:
    doc = extract_pdf(factory.stamped_scan_pdf(), ocr=OCR_OFF)
    assert doc.pages[0].note == "ocr_unavailable"
    assert doc.unreadable_pages == [1]
    assert any("could not be read" in w for w in doc.warnings)


@needs_tesseract
def test_d4_text_layer_is_kept_when_ocr_adds_nothing() -> None:
    doc = extract_pdf(factory.logo_cover_pdf(), ocr=OCR_ON)
    assert doc.pages[0].extraction == "text"
    assert doc.canonical_text == "Master Services Agreement - Cover Page"


def test_d4_text_rich_page_with_background_image_is_not_ocrd(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def must_not_run(*args, **kwargs):  # type: ignore[no-untyped-def]
        raise AssertionError("OCR must not run on a page with a full text layer")

    monkeypatch.setattr(pdf_extract, "ocr_page", must_not_run)
    assert (
        extract_pdf(factory.background_image_text_pdf(), ocr=OCR_ON).pages[0].extraction == "text"
    )


# --- D5: scanned page with no recognisable text labelled "blank" ----------------------------------


@needs_tesseract
def test_d5_image_page_without_recognisable_text_is_unreadable_not_blank() -> None:
    doc = extract_pdf(factory.picture_only_pdf(), ocr=OCR_ON)
    assert [(p.extraction, p.note) for p in doc.pages] == [("text", None), ("none", "ocr_no_text")]
    assert doc.unreadable_pages == [2]
    assert any("No text could be recognised on page(s) 2" in w for w in doc.warnings)


def test_d5_ink_detection_separates_blank_from_scanned_pages() -> None:
    blank = pymupdf.open(stream=factory.pdf_with_blank_page(), filetype="pdf")[1]
    scanned = pymupdf.open(stream=factory.scanned_pdf(1), filetype="pdf")[0]
    assert pdf_extract.ink_fraction(blank) < pdf_extract.MIN_INK_FRACTION
    assert pdf_extract.ink_fraction(scanned) >= pdf_extract.MIN_INK_FRACTION


def test_d5_genuinely_blank_page_is_still_blank_and_not_unreadable() -> None:
    doc = extract_pdf(factory.pdf_with_blank_page(), ocr=OCR_ON)
    assert doc.pages[1].note == "blank" and doc.unreadable_pages == []


# --- OCR progress (audit item 3: verified correct, now pinned by a test) ------------------------


@needs_tesseract
def test_ocr_progress_reports_each_page_in_order() -> None:
    events: list[tuple[float, str, bool]] = []
    extract_pdf(factory.scanned_pdf(pages=3), ocr=OCR_ON, on_progress=lambda *e: events.append(e))
    ocr_events = [(fraction, detail) for fraction, detail, is_ocr in events if is_ocr]
    assert [d for _, d in ocr_events] == ["OCR page 1 of 3", "OCR page 2 of 3", "OCR page 3 of 3"]
    fractions = [f for f, _ in ocr_events]
    assert fractions == sorted(fractions) and 0 < fractions[0] < fractions[-1] < 1
