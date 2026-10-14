"""Generate test documents in code, so tests never depend on files lying around.

Every generator returns bytes. PDFs come from PyMuPDF, DOCX files from python-docx.
"""

import io
import zipfile

import docx
import pymupdf

CLAUSE = (
    "The Company may terminate this Agreement upon thirty (30) days written notice to the "
    "Supplier. Termination shall not affect any rights accrued before the termination date."
)


def text_pdf(pages: int = 3, text: str = CLAUSE) -> bytes:
    document = pymupdf.open()
    for number in range(1, pages + 1):
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(72, 72, 520, 760), f"{number}. {text}", fontsize=11)
    return document.tobytes()


def _image_only_page(document: pymupdf.Document, text: str) -> None:
    """A page that *looks* like text but is only an image, like a scan."""
    source = pymupdf.open()
    source.new_page().insert_textbox(pymupdf.Rect(72, 72, 520, 760), text, fontsize=11)
    pixmap = source[0].get_pixmap(dpi=200)  # typical scanner resolution
    page = document.new_page()
    page.insert_image(page.rect, stream=pixmap.tobytes("jpeg", jpg_quality=80))  # like a scanner


def scanned_pdf(pages: int = 2) -> bytes:
    document = pymupdf.open()
    for _ in range(pages):
        _image_only_page(document, CLAUSE)
    return document.tobytes()


def mixed_pdf() -> bytes:
    """One text page followed by one scanned page."""
    document = pymupdf.open(stream=text_pdf(pages=1), filetype="pdf")
    _image_only_page(document, CLAUSE)
    return document.tobytes()


def encrypted_pdf(user_password: str = "secret") -> bytes:
    document = pymupdf.open(stream=text_pdf(pages=1), filetype="pdf")
    return document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_password, owner_pw="owner"
    )


def owner_restricted_pdf() -> bytes:
    """Encrypted with an owner password only: opens without a password (common for contracts)."""
    document = pymupdf.open(stream=text_pdf(pages=1), filetype="pdf")
    return document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", permissions=0)


def corrupt_pdf() -> bytes:
    """Cut off before any page objects: unrecoverable."""
    return text_pdf(pages=2)[:300]


def damaged_pdf() -> bytes:
    """Cut in half: MuPDF repairs it, recovering only some page content."""
    good = text_pdf(pages=2)
    return good[: len(good) // 2]


def docx_file(paragraphs: list[str] | None = None, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs if paragraphs is not None else [CLAUSE]:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def corrupt_docx() -> bytes:
    """A zip with the right part names but broken XML inside."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types")
        archive.writestr("word/document.xml", "<w:document><not-closed>")
    return buffer.getvalue()


def plain_zip() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    return buffer.getvalue()


OLE_HEADER = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504


_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _numbered(paragraph, num_id: int, level: int) -> None:  # type: ignore[no-untyped-def]
    from docx.oxml import parse_xml

    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(
        parse_xml(
            f'<w:numPr {_W_NS}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )


def numbered_docx() -> bytes:
    """A contract whose clause numbers exist only as Word auto-numbering (not in the text):
    1. / 1.1 / (a) levels, a heading-styled title, a table, and tracked changes."""
    from docx.oxml import parse_xml

    document = docx.Document()
    numbering = document.part.numbering_part.element
    numbering.append(
        parse_xml(
            f'<w:abstractNum {_W_NS} w:abstractNumId="90">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%1."/></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%1.%2"/></w:lvl>'
            '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
            '<w:lvlText w:val="(%3)"/></w:lvl>'
            "</w:abstractNum>"
        )
    )
    numbering.append(
        parse_xml(f'<w:num {_W_NS} w:numId="90"><w:abstractNumId w:val="90"/></w:num>')
    )

    document.add_heading("MASTER SERVICES AGREEMENT", level=1)
    document.add_paragraph("This Agreement is made between Acme Ltd and Beta LLC.")
    for text, level in [
        ("DEFINITIONS", 0),
        ("“Services” means the services described in Schedule 1.", 1),
        ("TERMINATION", 0),
        ("Either party may terminate this Agreement on thirty (30) days written notice.", 1),
        ("The Company may terminate immediately if the Supplier:", 1),
        ("breaches any material term; or", 2),
        ("becomes insolvent.", 2),
        ("LIABILITY", 0),
        ("The total liability of either party is capped at AED 100,000.", 1),
    ]:
        _numbered(document.add_paragraph(text), 90, level)

    tracked = document.add_paragraph("Payment is due within ")
    tracked._p.append(
        parse_xml(
            f'<w:ins {_W_NS} w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:t>forty-five</w:t></w:r></w:ins>"
        )
    )
    tracked._p.append(
        parse_xml(
            f'<w:del {_W_NS} w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:delText>thirty</w:delText></w:r></w:del>"
        )
    )
    tracked.add_run(" days.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Fee", "AED 100,000"
    table.cell(1, 0).text, table.cell(1, 1).text = "Term", "24 months"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def furniture_pdf(pages: int = 4) -> bytes:
    """Running header + 'Page n of N' footer on every page, and one sentence that continues
    across each page break."""
    document = pymupdf.open()
    for number in range(1, pages + 1):
        page = document.new_page()  # A4-ish 595 x 842
        page.insert_text((72, 40), "ACME CORP - CONFIDENTIAL", fontsize=9)
        page.insert_text((260, 815), f"Page {number} of {pages}", fontsize=9)
        start = "within ten (10) business days of the order. " if number > 1 else ""
        body = (
            f"{start}Clause {number} text: the parties agree to cooperate in good faith "
            "on all matters. The Supplier shall deliver the goods"
        )
        page.insert_textbox(pymupdf.Rect(72, 120, 520, 760), body, fontsize=11)
    return document.tobytes()


NEEDLE = "The Escrow Agent shall release the Deposit within seven (7) days of completion."


def large_pdf_with_needle(pages: int = 150, needle_page: int = 137) -> bytes:
    document = pymupdf.open()
    for number in range(1, pages + 1):
        text = f"{number}. {CLAUSE} {CLAUSE}"
        if number == needle_page:
            text += f" {NEEDLE}"
        document.new_page().insert_textbox(pymupdf.Rect(72, 72, 520, 760), text, fontsize=11)
    return document.tobytes()


def merged_clauses_pdf() -> bytes:
    """Two clauses in one text box: MuPDF returns them as a single block."""
    document = pymupdf.open()
    document.new_page().insert_textbox(
        pymupdf.Rect(72, 72, 520, 760),
        "8.1 The Supplier shall maintain insurance.\n8.2 The Company may audit the Supplier.",
        fontsize=11,
    )
    return document.tobytes()


def headings_pdf() -> bytes:
    """Unnumbered headings set in a larger font above body text."""
    document = pymupdf.open()
    page = document.new_page()
    y = 72
    for heading, body in [
        ("Payment Terms", "Invoices are payable within thirty days of receipt by the Customer."),
        ("Governing Law", "This Agreement is governed by the laws of England and Wales."),
    ]:
        page.insert_text((72, y), heading, fontsize=16)
        page.insert_textbox(pymupdf.Rect(72, y + 10, 520, y + 80), body, fontsize=10)
        y += 120
    return document.tobytes()


def pdf_with_blank_page() -> bytes:
    document = pymupdf.open(stream=text_pdf(pages=1), filetype="pdf")
    document.new_page()  # nothing on it: not scanned, simply blank
    return document.tobytes()


# --- Phase 3 audit regression fixtures ---------------------------------------------------------


def bold_headings_docx() -> bytes:
    """Headings that are just bold 'Normal' paragraphs (no Heading style), a bold cover block,
    a partly-bold label line and a bold lead-in ending with ':'."""
    document = docx.Document()

    def para(text: str, bold: bool | None = None) -> None:
        document.add_paragraph().add_run(text).bold = bold

    for cover in ["Personality and Human Values", "JOURNAL FOR SUCCESS", "Amity University"]:
        para(cover, bold=True)  # a title block: consecutive emphasised lines
    para("Personal Information", bold=True)
    label = document.add_paragraph()
    label.add_run("Name of the student: ").bold = True
    label.add_run("A. Student")  # only partly bold: not a heading
    para("About Yourself", bold=True)
    para("1. What are your strengths and weaknesses?")
    para("I am curious and like understanding how things work.")
    para("Activity 1: Learning Outcomes:", bold=True)  # lead-in ending with ':'
    para("I learned the value of planning a weekly timetable.")
    para("My Learning Experiences", bold=True)
    para("1. How did this course help you grow? Explain in the context of your strengths.")
    para("It strengthened my curiosity and teamwork.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _picture(width: int = 400, height: int = 500, *, with_text: bool = False) -> bytes:
    """A raster image: shapes only (no text), or a rendered paragraph of text."""
    source = pymupdf.open()
    page = source.new_page(width=width, height=height)
    if with_text:
        page.insert_textbox(pymupdf.Rect(30, 30, width - 30, height - 30), CLAUSE * 3, fontsize=11)
    else:
        for i in range(8):
            page.draw_circle((width / 2, height / 2), 20 + 22 * i, color=(0, 0, 0), width=3)
        page.draw_rect(pymupdf.Rect(40, 40, width - 40, height - 40), color=(0, 0, 0), width=6)
    return source[0].get_pixmap(dpi=200).tobytes("jpeg", jpg_quality=80)


def stamped_scan_pdf() -> bytes:
    """A scanned page (full-page image of text) that also carries a typed stamp in its text
    layer, so it has *some* extractable text: the stamp must not hide the unread scan."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, stream=_picture(595, 842, with_text=True))
    page.insert_text((300, 825), "CONFIDENTIAL  ACME-0000123  Page 1 of 1", fontsize=8)
    return document.tobytes()


def picture_only_pdf() -> bytes:
    """Page 1 text; page 2 a scanned diagram with no recognisable text."""
    document = pymupdf.open(stream=text_pdf(pages=1), filetype="pdf")
    page = document.new_page()
    page.insert_image(page.rect, stream=_picture())
    return document.tobytes()


def logo_cover_pdf() -> bytes:
    """A real text layer (short title) over a large non-text image: OCR adds nothing, so the
    accurate text layer must be kept."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(50, 50, 545, 700), stream=_picture())
    page.insert_text((72, 760), "Master Services Agreement - Cover Page", fontsize=14)
    return document.tobytes()


def background_image_text_pdf() -> bytes:
    """A full-page background image under a normal, text-rich page: no OCR needed."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, stream=_picture())
    page.insert_textbox(pymupdf.Rect(72, 72, 520, 760), CLAUSE * 4, fontsize=11)
    return document.tobytes()
